"""论文文本解析与视觉解析开关测试。"""

from scifig_mcp.tools import parse


def test_parse_paper_visual_disabled(monkeypatch, tmp_path):
    pdf = tmp_path / "paper.pdf"
    pdf.write_bytes(b"placeholder")
    monkeypatch.setattr(parse, "_parse_lightweight", lambda p: ("# paper", {"stem": "paper"}))

    def fail_if_called(*args, **kwargs):
        raise AssertionError("visual augmentation should be disabled")

    monkeypatch.setattr(parse, "_augment_with_visuals", fail_if_called)
    text, meta = parse.parse_paper_file(str(pdf), visual=False)

    assert text == "# paper"
    assert meta["visual"] == {"enabled": False, "status": "disabled"}


def test_parse_paper_appends_visual_analysis(monkeypatch, tmp_path):
    pdf = tmp_path / "paper.pdf"
    pdf.write_bytes(b"placeholder")
    monkeypatch.setattr(parse, "_parse_lightweight", lambda p: ("# paper", {"stem": "paper"}))
    monkeypatch.setattr(
        parse,
        "_augment_with_visuals",
        lambda p, text: (
            text + "\n\n## 图表与公式视觉解析\n",
            {"visual": {"enabled": True, "status": "ok", "pages": [1]}},
        ),
    )

    text, meta = parse.parse_paper_file(str(pdf), visual=True)
    assert "图表与公式视觉解析" in text
    assert meta["visual"]["status"] == "ok"


def test_parse_document_docx_extracts_text_and_tables(tmp_path):
    from docx import Document

    path = tmp_path / "paper.docx"
    doc = Document()
    doc.add_heading("方法", level=1)
    doc.add_paragraph("这是方法段落。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "准确率"
    table.cell(1, 1).text = "95%"
    doc.save(path)

    text, meta = parse.parse_document_file(str(path), visual=False)

    assert "# 方法" in text
    assert "这是方法段落" in text
    assert "准确率" in text and "95%" in text
    assert meta["document_type"] == "docx"
    assert meta["visual"]["status"] == "disabled"
